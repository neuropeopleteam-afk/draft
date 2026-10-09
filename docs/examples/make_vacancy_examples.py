"""Пример файла вакансии для проверки «Загрузить файл» на экране новой потребности (текст свой, компания вымышлена).

Собирает vacancy_python_middle.docx (python-docx — только для этого скрипта, в приложение не входит) и .txt;
PDF — конвертацией DOCX в LibreOffice:
    python3 docs/examples/make_vacancy_examples.py
    soffice --headless --convert-to pdf --outdir docs/examples docs/examples/vacancy_python_middle.docx
Копии кладутся в frontend/public/examples/ — ссылка «пример файла» рядом с кнопкой загрузки.
"""
from pathlib import Path
import shutil

TITLE = "Python-разработчик (middle) в команду тарифов"
ABOUT = ("Демо-компания «Северный склад» (вымышленная) делает сервис доставки для интернет-магазинов. Команда тарифов "
         "считает стоимость доставки для 40 тысяч заказов в час и отвечает за API для партнёров.")
TASKS = [
    "развивать сервис расчёта тарифов и API для партнёров;",
    "проектировать схемы данных и оптимизировать запросы к базе;",
    "писать тесты и разбирать инциденты вместе с дежурной сменой.",
]
MUST = [
    "Python 3 — коммерческий опыт от 3 лет;",
    "FastAPI или Django REST framework;",
    "PostgreSQL: индексы, планы запросов, транзакции;",
    "Docker и Git в ежедневной работе.",
]
PLUS = [
    "Redis для кэша и очередей;",
    "Kafka;",
    "опыт с Kubernetes.",
]
CONDITIONS = [
    ("Формат", "удалённо или гибрид (офис в Москве)"),
    ("Вилка", "230 000 – 290 000 руб. до вычета НДФЛ"),
    ("Занятость", "полная, оформление по ТК РФ"),
]
CONTACT = "Контактное лицо: рекрутер Анна Котова, hr@example.invalid, +7 900 000-00-00"


def as_text() -> str:
    lines = [TITLE, "", ABOUT, "", "Задачи:", *[f"- {x}" for x in TASKS], "", "Обязательно:", *[f"- {x}" for x in MUST],
             "", "Будет плюсом:", *[f"- {x}" for x in PLUS], "", "Условия:", *[f"{a}: {b}" for a, b in CONDITIONS],
             "", CONTACT]
    return "\n".join(lines) + "\n"


def main() -> None:
    from docx import Document  # системный python-docx: только для генерации примера
    from docx.oxml.ns import qn

    here = Path(__file__).resolve().parent
    doc = Document()
    for name in ("Normal", "Heading 1", "Heading 2", "List Bullet", "Table Grid"):
        st = doc.styles[name]
        st.font.name = "Arial"  # шрифт с кириллицей у всех: иначе при экспорте в PDF подставится шрифт без пробелов
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), "Arial")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attr), None)
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(ABOUT)
    for head, items in (("Задачи", TASKS), ("Обязательно", MUST), ("Будет плюсом", PLUS)):
        doc.add_heading(head, level=2)
        for x in items:
            doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("Условия", level=2)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for a, b in CONDITIONS:
        row = t.add_row().cells
        row[0].text, row[1].text = a, b
    doc.add_paragraph("")
    doc.add_paragraph(CONTACT)
    doc.core_properties.author = "Пример"
    doc.core_properties.title = TITLE
    doc.save(here / "vacancy_python_middle.docx")
    (here / "vacancy_python_middle.txt").write_text(as_text(), encoding="utf-8")
    pub = here.parent.parent / "frontend" / "public" / "examples"
    pub.mkdir(parents=True, exist_ok=True)
    for ext in ("docx", "txt", "pdf"):
        src = here / f"vacancy_python_middle.{ext}"
        if src.exists():
            shutil.copy(src, pub / src.name)


if __name__ == "__main__":
    main()
